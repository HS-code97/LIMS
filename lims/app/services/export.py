"""Excel 내보내기 (7장 P7, openpyxl), 각인 CSV·라벨 PDF (3.6절)."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Attachment, RunSpecimen, Specimen, TestResult, TestRun
from . import param_dict as pdict
from .actuals import badge_text
from .search import effective_map

HEAD_FILL = PatternFill("solid", fgColor="EEF0F2")
DEV_FILL = PatternFill("solid", fgColor="FDEBDD")


def _style_head(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def _autowidth(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 48)


def specimens_xlsx(db: Session, specs: list[Specimen], keys: list[str], basis: str, base_url: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "시편"
    keys = keys or pdict.quick_keys()
    head = ["약식번호", "정식번호", "단계", "상태", "과제", "공정런", "공정일", "원소재"]
    for k in keys:
        p = pdict.param(k)
        head += [f"{p['section_nm']} {p['l']} ({p['u']})", "출처"]
    head += ["시험런", "사진 수", "링크"]
    ws.append(head)
    emap = effective_map(db, [s.id for s in specs], keys)
    for s in specs:
        r = s.process_run
        row = [s.short_id, s.display_id, s.stage_name, s.status, s.project_code, r.run_no if r else "",
               r.run_date if r else None, s.material.coil_no if s.material else ""]
        dev_cols = []
        for k in keys:
            e = emap.get((s.id, k))
            v = None
            if e is not None:
                v = {"target": e.target, "actual": e.actual}.get(basis, e.effective)
                if e.out_of_tol:
                    dev_cols.append(len(row) + 1)
            row += [v, badge_text(e.source, e.level, e.inherited_from) if e else ""]
        runs = db.scalars(select(TestRun).join(RunSpecimen).where(RunSpecimen.specimen_id == s.id)).all()
        n_photo = db.query(Attachment).filter(Attachment.specimen_id == s.id, Attachment.is_deleted.is_(False)).count()
        row += [", ".join(t.test_run_no for t in runs), n_photo, f"{base_url}/specimens/{s.short_id}"]
        ws.append(row)
        for c in dev_cols:
            ws.cell(row=ws.max_row, column=c).fill = DEV_FILL
    _style_head(ws, len(head))
    _autowidth(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_run_xlsx(db: Session, tr: TestRun, base_url: str = "") -> bytes:
    wb = Workbook()
    tpl = tr.template
    info = wb.active
    info.title = "시험런"
    info.append(["시험런", tr.test_run_no])
    info.append(["템플릿", f"{tpl.code} v{tpl.version} {tpl.name}"])
    info.append(["설비", tr.equipment or ""])
    info.append(["시작", tr.start_date])
    info.append(["상태", tr.status])
    for c in tpl.condition_schema or []:
        info.append([f"{c['label']} ({c.get('unit', '')})", (tr.conditions or {}).get(c["key"])])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 40

    results = {(r.specimen_id, r.cycle_no): r for r in db.scalars(
        select(TestResult).where(TestResult.test_run_id == tr.id, TestResult.is_deleted.is_(False)))}
    photos: dict[int, list[Attachment]] = {}
    for a in db.scalars(select(Attachment).where(Attachment.result_id.in_([r.id for r in results.values()] or [0]),
                                                 Attachment.is_deleted.is_(False))):
        photos.setdefault(a.result_id, []).append(a)
    cycles = tr.cycles or tpl.default_cycles or 1
    for f in [x for x in tpl.field_schema if x.get("type") != "text"]:
        ws = wb.create_sheet(f["label"][:28])
        ws.append(["시편", "슬롯", "정식번호"] + [f"W{c:02d}" for c in range(cycles)])
        for b in sorted(tr.boardings, key=lambda b: b.specimen.short_id):
            row = [b.specimen.short_id, b.slot or "", b.specimen.display_id]
            for c in range(cycles):
                res = results.get((b.specimen_id, c))
                row.append((res.metrics or {}).get(f["key"]) if res else None)
            ws.append(row)
        _style_head(ws, 3 + cycles)
        _autowidth(ws)
    ws = wb.create_sheet("사진")
    ws.append(["시편", "회차", "파일", "링크"])
    for res in results.values():
        for a in photos.get(res.id, []):
            ws.append([res.specimen.short_id, f"W{res.cycle_no:02d}", a.original_name, f"{base_url}/files/{a.id}"])
    _style_head(ws, 4)
    _autowidth(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def engrave_csv(specs: list[Specimen]) -> bytes:
    """레이저 각인 CSV. 열·구분자·인코딩은 각인기 형식(11.3)에 맞춰 설정합니다.
    가능한 열: short_id, display_id, datamatrix_text(= 약식번호), condition_code, project_code, stage, serial_no"""
    from ..config import ENGRAVE_COLUMNS, ENGRAVE_DELIMITER, ENGRAVE_ENCODING, ENGRAVE_HEADER
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=ENGRAVE_DELIMITER, lineterminator="\r\n")
    cols = [c.strip() for c in ENGRAVE_COLUMNS if c.strip()]
    if ENGRAVE_HEADER:
        w.writerow(cols)
    for s in specs:
        vals = {"short_id": s.short_id, "display_id": s.engraved_label or s.display_id, "datamatrix_text": s.short_id,
                "condition_code": s.condition_code or "", "project_code": s.project_code, "stage": s.stage,
                "serial_no": s.serial_no}
        w.writerow([vals.get(c, "") for c in cols])
    return buf.getvalue().encode(ENGRAVE_ENCODING, errors="replace")


def labels_pdf(specs: list[Specimen], test_run: TestRun | None = None, slots: dict[int, str] | None = None) -> bytes:
    """라벨프린터용 PDF: 라벨 1장 = DataMatrix/QR(약식번호) + 정식번호 (+ 시험런·슬롯).
    크기·코드 종류는 라벨프린터 모델(11.3)에 맞춰 LIMS_LABEL_SIZE_MM, LIMS_LABEL_CODE로 설정."""
    from ..config import LABEL_CODE, LABEL_SIZE_MM
    from reportlab.graphics import renderPDF
    from reportlab.graphics.barcode import createBarcodeDrawing
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    W, H = LABEL_SIZE_MM[0] * mm, LABEL_SIZE_MM[1] * mm
    code_mm = min(H - 6 * mm, 14 * mm)
    tx = code_mm + 4.5 * mm
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(W, H))
    for s in specs:
        kind = "QR" if LABEL_CODE.lower() == "qr" else "ECC200DataMatrix"
        try:
            d = createBarcodeDrawing(kind, value=s.short_id)
        except Exception:
            d = createBarcodeDrawing("QR", value=s.short_id)
        scale = code_mm / max(d.width, d.height)
        d.scale(scale, scale)
        d.width *= scale
        d.height *= scale
        renderPDF.draw(d, c, 2.5 * mm, (H - code_mm) / 2)
        c.setFont("Courier-Bold", 9)
        c.drawString(tx, H - 7 * mm, s.short_id)
        c.setFont("Courier", 6.5)
        c.drawString(tx, H - 10.5 * mm, s.engraved_label or s.display_id)
        if test_run is not None:
            slot = (slots or {}).get(s.id) or ""
            c.drawString(tx, H - 15 * mm, f"{test_run.test_run_no}{' / ' + slot if slot else ''}")
        c.showPage()
    c.save()
    return buf.getvalue()
