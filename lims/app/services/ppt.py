"""PPT 자동 생성 (11.1 3단계, python-pptx). A4 가로(297×210 mm).

회사 표준 템플릿(11.3 확인 필요)은 LIMS_PPT_TEMPLATE에 .pptx 경로를 주면 그 파일의 마스터·레이아웃을 씁니다.
없으면 LIMS 디자인 토큰으로 만든 기본 서식을 씁니다.
"""
from __future__ import annotations

import io
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Mm, Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import FILES_DIR
from ..models import Attachment, Specimen, TestResult, TestRun
from . import actuals, lineage
from . import param_dict as pdict
from .templates_engine import cell_text
from .thumbnails import large_of

INK = RGBColor(0x17, 0x1C, 0x21)
INK3 = RGBColor(0x72, 0x77, 0x81)
PRIMARY = RGBColor(0x05, 0x52, 0x8E)
DEV = RGBColor(0xAC, 0x34, 0x00)
FONT = "맑은 고딕"
W, H = Mm(297), Mm(210)


def _prs() -> Presentation:
    tpl = os.getenv("LIMS_PPT_TEMPLATE")
    prs = Presentation(tpl) if tpl and Path(tpl).exists() else Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _blank(prs):
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    return prs.slides.add_slide(layout)


def _text(slide, x, y, w, h, text, size=12, bold=False, color=INK, mono=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        r.font.name = "Consolas" if mono else FONT
    return tb


def _header(slide, title, sub):
    _text(slide, Mm(12), Mm(8), Mm(270), Mm(12), title, 20, True, PRIMARY, mono=True)
    _text(slide, Mm(12), Mm(19), Mm(270), Mm(8), sub, 10, False, INK3)
    line = slide.shapes.add_connector(1, Mm(12), Mm(28), Mm(285), Mm(28))
    line.line.color.rgb = RGBColor(0xC9, 0xCE, 0xD6)


def _img(att: Attachment | None) -> io.BytesIO | None:
    """썸네일(WebP) → JPEG 메모리 스트림 (python-pptx는 WebP를 넣지 못함)."""
    if att is None:
        return None
    from PIL import Image
    for rel in (large_of(att.thumb_path), att.thumb_path):
        if rel and (FILES_DIR / rel).exists():
            with Image.open(FILES_DIR / rel) as im:
                buf = io.BytesIO()
                im.convert("RGB").save(buf, "JPEG", quality=85)
                buf.seek(0)
                return buf
    return None


def _table(slide, x, y, w, rows, col_w=None, size=9):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Mm(6) * len(rows))
    t = shp.table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = "" if val is None else str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.name, r.font.bold = Pt(size), FONT, ri == 0
    if col_w:
        for ci, cw in enumerate(col_w):
            t.columns[ci].width = cw
    return t


def test_run_pptx(db: Session, tr: TestRun) -> bytes:
    prs = _prs()
    tpl = tr.template
    results = {(r.specimen_id, r.cycle_no): r for r in db.scalars(
        select(TestResult).where(TestResult.test_run_id == tr.id, TestResult.is_deleted.is_(False)))}
    photos: dict[int, Attachment] = {}
    for a in db.scalars(select(Attachment).where(Attachment.result_id.in_([r.id for r in results.values()] or [0]),
                                                 Attachment.is_deleted.is_(False)).order_by(Attachment.id)):
        photos.setdefault(a.result_id, a)
    cycles = tr.cycles or tpl.default_cycles or 1
    boards = sorted(tr.boardings, key=lambda b: (b.slot or "~", b.specimen.short_id))

    s = _blank(prs)
    _header(s, tr.test_run_no, f"{tpl.name} · 템플릿 {tpl.code} v{tpl.version} · {tr.equipment or ''} · "
                               f"{tr.start_date} ~ {tr.end_date or '진행 중'}")
    cond = [f"{c['label']} {(tr.conditions or {}).get(c['key'], '')}{c.get('unit', '')}" for c in tpl.condition_schema or []]
    _text(s, Mm(12), Mm(32), Mm(270), Mm(10), " · ".join(cond) or "운전조건 없음", 11)
    gf = tpl.grid_fields[0] if tpl.grid_fields else None
    head = ["시편", "슬롯", "조건부"] + [f"W{c:02d}" for c in range(cycles)]
    rows = [head]
    for b in boards:
        rows.append([b.specimen.short_id, b.slot or "", b.specimen.condition_code or ""] +
                    [cell_text(tpl, results[(b.specimen_id, c)].metrics) if (b.specimen_id, c) in results else ""
                     for c in range(cycles)])
    _table(s, Mm(12), Mm(46), Mm(273), rows[:20], size=8)
    if gf:
        _text(s, Mm(12), Mm(196), Mm(270), Mm(8), f"표 값 = {gf['label']} ({gf.get('unit', '')}) · 빈칸 = 미입력", 8,
              color=INK3)

    for b in boards:
        s = _blank(prs)
        sp = b.specimen
        _header(s, sp.short_id, f"{sp.display_id} · {tr.test_run_no} · 슬롯 {b.slot or '-'}")
        n = min(cycles, 10)
        cols = 5
        cw, ch = Mm(53), Mm(40)
        for i in range(n):
            x, y = Mm(12) + (cw + Mm(2)) * (i % cols), Mm(34) + (ch + Mm(16)) * (i // cols)
            res = results.get((sp.id, i))
            img = _img(photos.get(res.id) if res else None)
            if img:
                s.shapes.add_picture(img, x, y, width=cw, height=ch)
            else:
                box = s.shapes.add_shape(1, x, y, cw, ch)
                box.fill.solid()
                box.fill.fore_color.rgb = RGBColor(0xEE, 0xF0, 0xF2)
                box.line.color.rgb = RGBColor(0xDE, 0xE3, 0xEA)
            _text(s, x, y + ch, cw, Mm(12), f"W{i:02d}  {cell_text(tpl, res.metrics) if res else ''}", 9, True, mono=True)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def specimen_pptx(db: Session, spec: Specimen) -> bytes:
    prs = _prs()
    s = _blank(prs)
    chain = " → ".join(n["label"] for n in lineage.chain(spec))
    _header(s, spec.short_id, f"{spec.display_id} · {spec.stage_name} · 계보 {chain}")
    rows = [["파라미터", "목표", "실측/대표값", "편차", "출처"]]
    if spec.process_run:
        eff = actuals.RunActuals(db, spec.process_run).all_for(spec)
        for k in pdict.quick_keys() + [k for k in eff if k not in pdict.quick_keys() and k != "out.cw"]:
            e = eff.get(k)
            if e is None:
                continue
            p = pdict.param(k)
            rows.append([f"{p['l']} ({p['u']})", f"{e.target:g}" if e.target is not None else "",
                         f"{e.effective:g}", (f"{e.deviation:+g}" + (" ●" if e.out_of_tol else "")) if e.deviation is not None else "",
                         e.badge])
    _table(s, Mm(12), Mm(34), Mm(150), rows[:24], [Mm(52), Mm(22), Mm(28), Mm(24), Mm(24)], size=8)
    ids = [n.id for n in [*lineage.ancestors(spec), spec]]
    atts = [a for a in db.scalars(select(Attachment).where(Attachment.specimen_id.in_(ids), Attachment.is_deleted.is_(False))
                                  .order_by(Attachment.uploaded_at.desc(), Attachment.id.desc())) if a.thumb_path][:6]
    for i, a in enumerate(atts):
        x, y = Mm(168) + Mm(59) * (i % 2), Mm(34) + Mm(54) * (i // 2)
        img = _img(a)
        if img:
            s.shapes.add_picture(img, x, y, width=Mm(57), height=Mm(43))
        _text(s, x, y + Mm(43), Mm(57), Mm(8), a.original_name[:38], 7, color=INK3, mono=True)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
