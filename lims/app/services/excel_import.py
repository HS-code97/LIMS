"""과거 엑셀 데이터 일괄 가져오기 (11.1 2단계).

양식(시트 이름 고정, 없는 시트는 건너뜀) — 처리 순서: 공정런 → 분기 → 실측 → 시험결과
  공정런:   키 | 과제 | 일자 | 프리셋 | 목표부착량 | 원판수 | 시작번호 | 원소재 | 메모 | (사전 키 열: pot.T, rth.DP …)
            '키'는 이 파일 안에서만 쓰는 이름(예: R1). 실측 시트에서 공정런을 가리킬 때 씁니다.
  분기:     부모 | 단계 | 개수 | 시작번호 | 크기 | 위치 | 방향 | 평가면
  실측:     공정런(키 또는 PR-번호) | 시편(빈칸 = 세션 공통) | (사전 키 또는 라벨 열…) | 방법
  시험결과: 그룹 | 템플릿 | 설비 | 시작일 | 시편 | 슬롯 | 회차 | (템플릿 필드 키 또는 라벨 열…)
            같은 '그룹' 값의 행은 한 시험런으로 묶입니다(시험런 번호는 자동 발번).
모든 시트를 한 트랜잭션으로 처리하고, 오류가 하나라도 있으면 아무것도 저장하지 않습니다.
"""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Material, ProcessPreset, ProcessRun, TestTemplate
from ..schemas import ProcessRunIn, SplitIn, TestRunIn
from . import actuals, workflow
from . import param_dict as pdict
from .templates_engine import MetricsError

SHEETS = ("공정런", "분기", "실측", "시험결과")


def _rows(ws) -> list[dict]:
    it = ws.iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        return []
    head = [str(h).strip() if h is not None else "" for h in head]
    out = []
    for r in it:
        if r is None or all(v in (None, "") for v in r):
            continue
        out.append({h: v for h, v in zip(head, r) if h})
    return out


def _date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v).strip()[:10])


def _keymap() -> dict[str, str]:
    km = {}
    for k, p in pdict.params().items():
        for name in (k, p["l"], f"{p['section_nm']} {p['l']}", *p.get("aka", [])):
            km[name.upper()] = k
    return km


def run_import(db: Session, data: bytes, user_id: int | None, dry_run: bool = True) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    report: dict = {"created": {"process_runs": [], "specimens": [], "actuals": 0, "test_runs": [], "results": 0},
                    "errors": [], "warnings": []}
    err = report["errors"]
    run_keys: dict[str, ProcessRun] = {}
    km = _keymap()
    presets = {p.name.upper(): p for p in db.scalars(select(ProcessPreset).where(ProcessPreset.is_active.is_(True)))}

    # ── 공정런 ──
    if "공정런" in wb.sheetnames:
        for i, r in enumerate(_rows(wb["공정런"]), start=2):
            where = f"공정런 {i}행"
            try:
                preset = presets.get(str(r.get("프리셋", "")).strip().upper())
                if preset is None:
                    err.append(f"{where}: 프리셋 '{r.get('프리셋')}'이 없습니다")
                    continue
                recipe = pdict.preset_recipe(preset.overrides)
                for col, v in r.items():
                    k = km.get(col.upper())
                    if k and v not in (None, "") and k != "out.cw":
                        sec, pk = k.split(".")
                        recipe.setdefault(sec, {})[pk] = float(v) if pdict.is_numeric(k) else v
                mat = None
                if r.get("원소재"):
                    mat = db.scalar(select(Material).where(Material.coil_no == str(r["원소재"]).strip().upper()))
                    if mat is None:
                        mat = Material(coil_no=str(r["원소재"]).strip().upper(), note="엑셀 가져오기")
                        db.add(mat)
                        db.flush()
                payload = ProcessRunIn(dict_version=pdict.version(), preset_id=preset.id,
                                       project_code=str(r.get("과제", "")).strip(), material_id=mat.id if mat else None,
                                       run_date=_date(r.get("일자")), coating_weight_target=float(r.get("목표부착량")),
                                       panel_count=int(r.get("원판수", 1)), serial_start=int(r.get("시작번호")),
                                       recipe=recipe, note=str(r.get("메모") or "엑셀 가져오기"),
                                       source="import", confirm_out_of_range=True)
                run = workflow.create_process_run(db, payload, user_id)
                run_keys[str(r.get("키") or run.run_no)] = run
                report["created"]["process_runs"].append(run.run_no)
                report["created"]["specimens"] += [p.short_id for p in run.panels]
            except workflow.Conflict as e:
                err.append(f"{where}: {e} ({', '.join(e.items)})")
            except (workflow.Invalid, ValueError, TypeError) as e:
                err.append(f"{where}: {e}")

    # ── 분기 ──
    if "분기" in wb.sheetnames:
        for i, r in enumerate(_rows(wb["분기"]), start=2):
            where = f"분기 {i}행"
            parent = workflow.find_specimen(db, str(r.get("부모", "")))
            if parent is None:
                err.append(f"{where}: 부모 시편 '{r.get('부모')}'이 없습니다")
                continue
            try:
                kids = workflow.split(db, parent, SplitIn(
                    stage=str(r.get("단계", "C")).strip().upper(), count=int(r.get("개수", 1)),
                    serial_start=int(r.get("시작번호")), size_text=r.get("크기"), sample_position=r.get("위치"),
                    rolling_direction=r.get("방향"), surface_side=r.get("평가면")), user_id)
                report["created"]["specimens"] += [k.short_id for k in kids]
            except workflow.Conflict as e:
                err.append(f"{where}: {e} ({', '.join(e.items)})")
            except (workflow.Invalid, ValueError, TypeError) as e:
                err.append(f"{where}: {e}")

    # ── 실측 ──
    if "실측" in wb.sheetnames:
        touched = set()
        for i, r in enumerate(_rows(wb["실측"]), start=2):
            where = f"실측 {i}행"
            ref = str(r.get("공정런", "")).strip()
            run = run_keys.get(ref) or db.scalar(select(ProcessRun).where(ProcessRun.run_no == ref.upper()))
            if run is None:
                err.append(f"{where}: 공정런 '{ref}'을 찾지 못했습니다")
                continue
            spec = None
            if r.get("시편"):
                spec = workflow.find_specimen(db, str(r["시편"]))
                if spec is None or spec.process_run_id != run.id:
                    err.append(f"{where}: 시편 '{r['시편']}'이 {run.run_no}에 없습니다")
                    continue
            for col, v in r.items():
                if col in ("공정런", "시편", "방법") or v in (None, ""):
                    continue
                k = km.get(col.upper())
                if k is None:
                    report["warnings"].append(f"{where}: 사전에 없는 열 '{col}' 무시")
                    continue
                actuals.add_actual(db, run, k, v, specimen=spec, source="manual", method=r.get("방법"),
                                   user_id=user_id, note="엑셀 가져오기")
                report["created"]["actuals"] += 1
            touched.add(run.id)
        for run in {r for r in run_keys.values()} | set(db.scalars(select(ProcessRun).where(ProcessRun.id.in_(touched)))):
            actuals.recompute_run(db, run)

    # ── 시험결과 ──
    if "시험결과" in wb.sheetnames:
        groups: dict[str, list[tuple[int, dict]]] = {}
        for i, r in enumerate(_rows(wb["시험결과"]), start=2):
            groups.setdefault(str(r.get("그룹") or f"행{i}"), []).append((i, r))
        for g, rows in groups.items():
            first = rows[0][1]
            tpl = db.scalar(select(TestTemplate).where(TestTemplate.code == str(first.get("템플릿", "")).strip().upper(),
                                                       TestTemplate.is_active.is_(True)))
            if tpl is None:
                err.append(f"시험결과 그룹 {g}: 템플릿 '{first.get('템플릿')}'이 없습니다")
                continue
            fmap = {f["key"].upper(): f["key"] for f in tpl.field_schema} | \
                   {f["label"].upper(): f["key"] for f in tpl.field_schema}
            boards, missing = {}, []
            for _, r in rows:
                sid = str(r.get("시편", "")).strip().upper()
                if workflow.find_specimen(db, sid) is None:
                    missing.append(sid)
                elif sid not in boards:
                    boards[sid] = r.get("슬롯")
            if missing:
                err.append(f"시험결과 그룹 {g}: 없는 시편 {', '.join(sorted(set(missing)))}")
                continue
            try:
                tr = workflow.create_test_run(db, TestRunIn(
                    template_id=tpl.id, equipment=first.get("설비"), start_date=_date(first.get("시작일")),
                    specimens=[{"short_id": s, "slot": str(sl) if sl else None} for s, sl in boards.items()],
                    note="엑셀 가져오기"), user_id)
                report["created"]["test_runs"].append(tr.test_run_no)
                for i, r in rows:
                    metrics = {fmap[c.upper()]: v for c, v in r.items() if c.upper() in fmap and v not in (None, "")}
                    if not metrics:
                        continue
                    spec = workflow.find_specimen(db, str(r["시편"]))
                    _, w = workflow.save_result(db, tr, spec, int(r.get("회차") or 0), metrics, user_id)
                    report["warnings"] += [f"시험결과 {i}행: {x}" for x in w]
                    report["created"]["results"] += 1
            except MetricsError as e:
                err.append(f"시험결과 그룹 {g}: {'; '.join(e.errors)}")
            except (workflow.Invalid, ValueError, TypeError) as e:
                err.append(f"시험결과 그룹 {g}: {e}")

    if err or dry_run:
        db.rollback()
    else:
        db.commit()
    report["dry_run"] = dry_run
    report["saved"] = not err and not dry_run
    return report


def template_workbook() -> bytes:
    """빈 양식 + 예시 1행."""
    wb = Workbook()
    ws = wb.active
    ws.title = "공정런"
    ws.append(["키", "과제", "일자", "프리셋", "목표부착량", "원판수", "시작번호", "원소재", "메모", "pot.T", "rth.DP"])
    ws.append(["R1", "26P4", "2025-03-14", "M2", 60, 3, 1, "POS-25-1234", "과거 데이터", 590, -40])
    ws = wb.create_sheet("분기")
    ws.append(["부모", "단계", "개수", "시작번호", "크기", "위치", "방향", "평가면"])
    ws.append(["26P4-P0001", "C", 2, 1, "70x150", "C", "RD", "TOP"])
    ws = wb.create_sheet("실측")
    ws.append(["공정런", "시편", "방법", "욕온", "이슬점", "최고 판온", "부착량 윗면", "부착량 아랫면"])
    ws.append(["R1", "", "열전대", 588, -38, None, None, None])
    ws.append(["R1", "26P4-P0001", "중량법", None, None, 801, 57, 59])
    ws = wb.create_sheet("시험결과")
    ws.append(["그룹", "템플릿", "설비", "시작일", "시편", "슬롯", "회차", "red_rust_pct", "white_rust_pct"])
    ws.append(["G1", "CCT_STD", "CCT-01", "2025-03-20", "26P4-C0001", "A-01", 0, 0, 0])
    ws.append(["G1", "CCT_STD", "CCT-01", "2025-03-20", "26P4-C0001", "A-01", 1, 0, 5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
